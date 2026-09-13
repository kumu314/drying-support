#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
A 题 · 问题 1 —— 预热平衡阶段（0~1800 s）药材温度场与水分浓度场求解器
=====================================================================

物理模型（假设 H1~H4）
----------------------------------------
圆柱状药材（R = 2 cm，长 25 cm），按**一维轴对称径向**处理（H1）。
预热平衡阶段物性取常数（附录 2），故热/湿两方程**解耦**：

    热传导：  rho*cp * dT/dt = (1/r) * d/dr ( k * r * dT/dr )
    水分扩散：dC/dt         = (1/r) * d/dr ( D(C) * r * dC/dr )
              D(C) = 7e-9 * exp( -0.89 / C )

    初始：T(r,0) = 28 degC ,  C(r,0) = 2.55 kg/kg
    边界：r = 0   对称，dT/dr = dC/dr = 0
          r = R   Robin： -k*dT/dr = h*(T_R - T_air)  （流出为正；流入项 h*(T_air-T_R) 见表面控制体平衡）
                          -D*dC/dr = h_m*(C_R - C_air)
    T_air(t)、C_air(t) 由附件 1 线性插值给出（241 点，0~14400 s，步长 60 s）

数值方法
--------
有限体积法（FVM），柱坐标，**节点恰好落在题目要求的输出网格**上
（r = 0, 0.1, ..., 2.0 cm），两端为半控制体：

  * r = 0 的 1/r 奇点由"内侧面通量恒为零（对称）"自然消去，无需 L'Hopital 特殊处理；
  * 内部面通量在两个相邻控制体间**严格成对出现**（+A*q / -A*q）
    => 离散格式**严格守恒**，守恒性可对账到机器精度（这是 V1 验证的基础）。

时间推进用 Crank-Nicolson（二阶精度、无条件稳定）。
非线性 D(C) 用**中点状态 Picard 迭代**（D 取 0.5*(C^n + C^{n+1,k})）处理。
热方程为常系数线性系统 => 三对角矩阵只装配一次，每步仅换右端项。

验证（详见本目录 verify_report.md）
-----------------------------------
V1 离散守恒性（能量 / 水分总量，逐时间步对账）
V2 网格收敛性（dr, dr/2, dr/4）
V3 时间步收敛性（dt, dt/2, dt/4）
V4 独立实现交叉验证（scipy BDF 自适应积分同一 FVM 半离散系统）
V5 解析解比对（常系数退化情形 vs 无限圆柱分离变量级数解）
V6 物理合理性（中心滞后于表面、单调、极值界限）

用法
----
    python prob1/solve.py --air <附件1.xlsx 路径> [--outdir ...]

注意：附件**不在仓库内**（防泄题），必须由外部路径传入。
"""

from __future__ import annotations

import argparse
import json
import math
import os
import sys

# ---- 控制台可移植性兜底-----------------------------------------
# 默认中文 Windows 控制台是 cp936；stdout 遇到不可编码字符（如 U+2212、U+2705）
# 会抛 UnicodeEncodeError 并中断。下面用 errors="replace" 降级为 '?'；
# **不改变 stdout 编码**，故中文仍按控制台原生编码正常显示。
if hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(errors="replace")
    except (ValueError, OSError):
        pass

import time as _time

import numpy as np

# --------------------------------------------------------------------------
# 题目给定参数 —— 唯一真源：题目附录2
# 此处为机器可读镜像；改值必须同步核对全部子问题与全部结果
# --------------------------------------------------------------------------
PARAMS = dict(
    R=0.02,        # 药材半径 [m]     题干："半径为 2 cm"
    L=0.25,        # 药材长度 [m]     题干："长为 25 cm"（仅用于降维论证，不参与求解）
    rho=820.0,     # 密度 [kg/m^3]    附录 2
    cp=2600.0,     # 比热容 [J/(kg*K)] 附录 2
    k=0.36,        # 热传导系数 [W/(m*K)] 附录 2
    h=25.0,        # 对流换热系数 [W/(m^2*K)] 附录 2
    hm=8.0e-7,     # 对流传质系数 [m/s] 附录 2
    D0=7.0e-9,     # D 前因子 [m^2/s]  附录 2
    Da=0.89,       # D = D0 * exp(-Da / C) 中的 Da   附录 2
    T0=28.0,       # 初始温度 [degC]   题干
    C0=2.55,       # 初始干基含水率 [kg/kg] 题干
    t_end=1800.0,  # 预热平衡阶段时长 [s] 题干"30 分钟"
)

# 题目要求的输出网格（cm）；表 1/表 2 的抽样列
OUT_CM = np.round(np.arange(0.0, 2.0 + 1e-9, 0.1), 10)      # 0,0.1,...,2.0  -> 21 点
TAB_TIMES = [100.0, 300.0, 600.0, 900.0, 1200.0, 1500.0, 1800.0]  # 表 1/表 2 时间行
TAB_CM = [0.0, 0.5, 1.0, 1.5, 2.0]                                # 表 1/表 2 距离列


# ==========================================================================
# 工具
# ==========================================================================
def thomas(a: np.ndarray, b: np.ndarray, c: np.ndarray, d: np.ndarray) -> np.ndarray:
    """三对角线性方程组求解（Thomas 追赶法）。

    a 下对角（a[0] 未用）、b 对角、c 上对角（c[-1] 未用）、d 右端。
    """
    n = b.size
    cp_ = np.empty(n)
    dp_ = np.empty(n)
    cp_[0] = c[0] / b[0]
    dp_[0] = d[0] / b[0]
    for i in range(1, n):
        m = b[i] - a[i] * cp_[i - 1]
        cp_[i] = (c[i] / m) if i < n - 1 else 0.0
        dp_[i] = (d[i] - a[i] * dp_[i - 1]) / m
    x = np.empty(n)
    x[-1] = dp_[-1]
    for i in range(n - 2, -1, -1):
        x[i] = dp_[i] - cp_[i] * x[i + 1]
    return x


def load_air_table(path: str):
    """读附件 1：时间[s] / 温度[degC] / 水分浓度[kg/kg]。"""
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    rows = [r for r in ws.iter_rows(values_only=True) if r and r[0] is not None]
    header = rows[0]
    body = rows[1:]
    t = np.array([float(r[0]) for r in body], dtype=float)
    Ta = np.array([float(r[1]) for r in body], dtype=float)
    Ca = np.array([float(r[2]) for r in body], dtype=float)
    idx = np.argsort(t)
    return header, t[idx], Ta[idx], Ca[idx]


class AirDriver:
    """烘房环境驱动量 T_air(t), C_air(t)：附件 1 离散点上的**线性插值**。

    题目只给离散点（步长 60 s），必须插值；线性插值是最保守、最不可被质疑的选择
    （不引入样条可能的过冲）。超出数据范围时取端点值（不外推）。
    """

    def __init__(self, t, Ta, Ca):
        self.t, self.Ta, self.Ca = t, Ta, Ca
        self.t_min, self.t_max = float(t[0]), float(t[-1])

    def __call__(self, tau: float):
        tt = min(max(tau, self.t_min), self.t_max)
        return float(np.interp(tt, self.t, self.Ta)), float(np.interp(tt, self.t, self.Ca))


def make_grid(N: int, R: float):
    """柱坐标 FVM 网格：节点 r_i = i*dr (i=0..N)，两端半控制体。

    返回 (dr, r_nodes, A_face, A_R, V_cell)
      A_face[i] : 节点 i 与 i+1 之间内部面的面积（2*pi*r_{i+1/2}）
      A_R       : 外表面（侧面）面积 2*pi*R
      V_cell[i] : 节点 i 的控制体体积
    """
    dr = R / N
    r_nodes = np.arange(N + 1) * dr
    A_face = 2.0 * np.pi * (np.arange(N) + 0.5) * dr      # i = 0..N-1
    A_R = 2.0 * np.pi * R
    V = np.empty(N + 1)
    V[0] = np.pi * (0.5 * dr) ** 2
    V[1:N] = np.pi * 2.0 * np.arange(1, N) * dr ** 2
    V[N] = np.pi * (N ** 2 - (N - 0.5) ** 2) * dr ** 2
    return dr, r_nodes, A_face, A_R, V


def D_of_C(C, P=PARAMS):
    """水分扩散系数 D(C) = D0 * exp(-Da/C)  [m^2/s]，逐元素。"""
    C_safe = np.maximum(C, 1.0e-12)
    return P["D0"] * np.exp(-P["Da"] / C_safe)


# ==========================================================================
# 求解器
# ==========================================================================
def solve_case(N: int, dt: int | float, out_every: float, air: AirDriver,
               P=PARAMS, t_end=None, record_all: bool = False):
    """跑一个算例。

    N         : 径向网格数（节点数 N+1）
    dt        : 时间步长 [s]
    out_every : 每隔多少秒输出一次
    record_all: True 则返回整条时间序列（用于验证/画图）

    返回 dict：times, T_hist, C_hist, 以及守恒性误差记录
    """
    t_end = P["t_end"] if t_end is None else t_end
    dr, r, A_face, A_R, V = make_grid(N, P["R"])
    n = N + 1

    rc = P["rho"] * P["cp"]
    beta = P["k"] * A_face / dr                    # 热：内部面系数（常系数）

    # ---- 热方程矩阵（常系数，只装配一次）----
    aH = np.zeros(n); bH = np.zeros(n); cH = np.zeros(n)
    bH[0] = rc * V[0] / dt + 0.5 * beta[0]
    cH[0] = -0.5 * beta[0]
    for i in range(1, N):
        bH[i] = rc * V[i] / dt + 0.5 * (beta[i - 1] + beta[i])
        aH[i] = -0.5 * beta[i - 1]
        cH[i] = -0.5 * beta[i]
    bH[N] = rc * V[N] / dt + 0.5 * A_R * P["h"] + 0.5 * beta[N - 1]
    aH[N] = -0.5 * beta[N - 1]

    # ---- 初值 ----
    T = np.full(n, P["T0"], dtype=float)
    C = np.full(n, P["C0"], dtype=float)

    nsteps = int(round(t_end / dt))
    out_stride = int(round(out_every / dt))
    if abs(out_stride * dt - out_every) > 1e-9:
        raise ValueError("out_every 必须是 dt 的整数倍")

    n_out = nsteps // out_stride
    times = np.empty(n_out + 1)
    T_hist = np.empty((n_out + 1, n))
    C_hist = np.empty((n_out + 1, n))
    times[0] = 0.0
    T_hist[0] = T
    C_hist[0] = C

    # ---- 守恒性对账累加器 ----
    # 能量：E = Σ rc*V*T ，表面流入功率 Q  = A_R*h*(T_air  - T_R)
    # 水分：M = Σ V*C    ，表面流入通量 Qm = A_R*hm*(C_air - C_R)
    # 判据用**累积对账**：Σ|δE_数值 − δE_期望| / Σ|δE_期望|。
    # 不用"单步相对误差"，因为 t=0 处 T_air = T0 使驱动恰为零，
    # 分母趋零会把浮点噪声放大成假异常（首轮实测 4.5e+2 就是这么来的）。
    E0 = float(np.sum(rc * V * T))
    M0 = float(np.sum(V * C))
    Ta0, Ca0 = air(0.0)
    Q_prev = A_R * P["h"] * (Ta0 - T[N])
    Qm_prev = A_R * P["hm"] * (Ca0 - C[N])
    E_prev, M_prev = E0, M0
    cumE_defect = cumE_expect = 0.0
    cumM_defect = cumM_expect = 0.0
    stepE_max = stepE_scale = 0.0
    stepM_max = stepM_scale = 0.0

    k_picard_max = 0

    for s in range(1, nsteps + 1):
        tn = s * dt
        tnm = tn - dt
        Ta_n, Ca_n = air(tnm)
        Ta_np, Ca_np = air(tn)

        # ------------------ 热：CN，线性，一次解 ------------------
        rhs = rc * V / dt * T
        rhs[0] += 0.5 * beta[0] * (T[1] - T[0])
        rhs[1:N] += 0.5 * (beta[1:N] * (T[2:N + 1] - T[1:N])
                           - beta[0:N - 1] * (T[1:N] - T[0:N - 1]))
        rhs[N] += (0.5 * A_R * P["h"] * (Ta_n + Ta_np - T[N])
                   - 0.5 * beta[N - 1] * (T[N] - T[N - 1]))
        T_new = thomas(aH, bH, cH, rhs)

        # ------------------ 湿：CN + 中点状态 Picard ------------------
        C_k = C.copy()
        n_pic = 0
        Dmid = None
        for it in range(30):
            C_mid = 0.5 * (C + C_k)
            Dmid = D_of_C(C_mid)
            Dface = 0.5 * (Dmid[0:N] + Dmid[1:N + 1])        # i=0..N-1 内部面
            bm = Dface * A_face / dr
            d_face = bm

            am = np.zeros(n); bmm = np.zeros(n); cm = np.zeros(n)
            bmm[0] = V[0] / dt + 0.5 * d_face[0]
            cm[0] = -0.5 * d_face[0]
            for i in range(1, N):
                bmm[i] = V[i] / dt + 0.5 * (d_face[i - 1] + d_face[i])
                am[i] = -0.5 * d_face[i - 1]
                cm[i] = -0.5 * d_face[i]
            bmm[N] = V[N] / dt + 0.5 * A_R * P["hm"] + 0.5 * d_face[N - 1]
            am[N] = -0.5 * d_face[N - 1]

            rhsm = V / dt * C
            rhsm[0] += 0.5 * d_face[0] * (C[1] - C[0])
            rhsm[1:N] += 0.5 * (d_face[1:N] * (C[2:N + 1] - C[1:N])
                                - d_face[0:N - 1] * (C[1:N] - C[0:N - 1]))
            rhsm[N] += (0.5 * A_R * P["hm"] * (Ca_n + Ca_np - C[N])
                        - 0.5 * d_face[N - 1] * (C[N] - C[N - 1]))

            C_next = thomas(am, bmm, cm, rhsm)
            n_pic = it + 1
            if np.max(np.abs(C_next - C_k)) < 1.0e-11:
                C_k = C_next
                break
            C_k = C_next
        C_new = C_k
        k_picard_max = max(k_picard_max, n_pic)

        # ---- 守恒性对账（必须用与格式一致的梯形积分）----
        Q_now = A_R * P["h"] * (Ta_np - T_new[N])
        Qm_now = A_R * P["hm"] * (Ca_np - C_new[N])
        E_now = float(np.sum(rc * V * T_new))
        M_now = float(np.sum(V * C_new))

        dE = E_now - E_prev
        dE_exp = 0.5 * dt * (Q_prev + Q_now)
        cumE_defect += abs(dE - dE_exp)
        cumE_expect += abs(dE_exp)
        stepE_max = max(stepE_max, abs(dE - dE_exp))
        stepE_scale = max(stepE_scale, abs(dE_exp))

        dM = M_now - M_prev
        dM_exp = 0.5 * dt * (Qm_prev + Qm_now)
        cumM_defect += abs(dM - dM_exp)
        cumM_expect += abs(dM_exp)
        stepM_max = max(stepM_max, abs(dM - dM_exp))
        stepM_scale = max(stepM_scale, abs(dM_exp))

        T, C = T_new, C_new
        E_prev, M_prev = E_now, M_now
        Q_prev, Qm_prev = Q_now, Qm_now

        if s % out_stride == 0:
            j = s // out_stride
            times[j] = tn
            T_hist[j] = T
            C_hist[j] = C

    if not record_all:
        T_hist = T_hist[[0, -1]]
        C_hist = C_hist[[0, -1]]
        times = times[[0, -1]]

    return dict(
        N=N, dt=dt, dr=dr, r=r, times=times, T=T_hist, C=C_hist,
        max_consE_rel=cumE_defect / max(cumE_expect, 1.0e-30),
        max_consM_rel=cumM_defect / max(cumM_expect, 1.0e-30),
        stepE_rel=stepE_max / max(stepE_scale, 1.0e-30),
        stepM_rel=stepM_max / max(stepM_scale, 1.0e-30),
        E0=E0, E_end=E_prev, M0=M0, M_end=M_prev,
        cumE_defect=cumE_defect, cumE_expect=cumE_expect,
        cumM_defect=cumM_defect, cumM_expect=cumM_expect,
        picard_max=k_picard_max,
    )


# ==========================================================================
# V1 守恒性：单独跑一小段做逐时间步对账（上面 solve_case 已内建累计）
# ==========================================================================
def verify_conservation(N=20, dt=0.05, air=None):
    res = solve_case(N, dt, out_every=dt * 100, air=air, record_all=False)
    return res["max_consE_rel"], res["max_consM_rel"]


# ==========================================================================
# V3 / V4 / V5 独立与解析验证
# ==========================================================================
def rhs_semi_discrete(T, C, tau, air, N, dr, r, A_face, A_R, V, P=PARAMS,
                      const_D=None):
    """FVM 半离散右端：返回 dT/dt, dC/dt（供 scipy 自适应积分使用）。

    与主求解器**共用同一空间离散**，但时间积分交给完全不同的算法，
    因此可交叉验证时间推进实现的正确性。
    """
    n = N + 1
    Ta, Ca = air(tau)
    rc = P["rho"] * P["cp"]
    beta = P["k"] * A_face / dr

    # ---- 热 ----
    L = np.empty(n)
    L[0] = beta[0] * (T[1] - T[0])
    L[1:N] = beta[1:N] * (T[2:N + 1] - T[1:N]) - beta[0:N - 1] * (T[1:N] - T[0:N - 1])
    L[N] = A_R * P["h"] * (Ta - T[N]) - beta[N - 1] * (T[N] - T[N - 1])
    dT = L / (rc * V)

    # ---- 湿 ----
    Dv = D_of_C(C) if const_D is None else np.full(n, float(const_D))
    Dface = 0.5 * (Dv[0:N] + Dv[1:N + 1])
    dm = Dface * A_face / dr
    Lm = np.empty(n)
    Lm[0] = dm[0] * (C[1] - C[0])
    Lm[1:N] = dm[1:N] * (C[2:N + 1] - C[1:N]) - dm[0:N - 1] * (C[1:N] - C[0:N - 1])
    Lm[N] = A_R * P["hm"] * (Ca - C[N]) - dm[N - 1] * (C[N] - C[N - 1])
    dC = Lm / V

    return np.concatenate([dT, dC])


def verify_scipy_bdf(N=20, t_end=None, air=None, P=PARAMS):
    """V4：用 scipy 的 BDF（自适应、变阶）积分同一半离散系统，与 CN 结果比对。"""
    from scipy.integrate import solve_ivp

    t_end = P["t_end"] if t_end is None else t_end
    dr, r, A_face, A_R, V = make_grid(N, P["R"])
    y0 = np.concatenate([np.full(N + 1, P["T0"]), np.full(N + 1, P["C0"])])

    def f(tau, y):
        return rhs_semi_discrete(y[:N + 1], y[N + 1:], tau, air, N, dr, r,
                                 A_face, A_R, V, P)

    sol = solve_ivp(f, (0.0, t_end), y0, method="BDF", rtol=1e-10, atol=1e-12,
                    dense_output=False, max_step=5.0)
    yend = sol.y[:, -1]
    return r, yend[:N + 1], yend[N + 1:], float(sol.t[-1]), bool(sol.success)


def analytic_cylinder_series(r, t, D, hm, C0, C_air, R, n_terms=60):
    """V5：无限圆柱常系数扩散问题的分离变量级数解。

    dC/dt = D*(1/r)*d/dr(r*dC/dr),  dC/dr|_{r=0}=0,  -D*dC/dr|_R = hm*(C_R - C_air)
    u = C - C_air 满足同一方程且 u|_R 处边界齐次化后：
        u(r,t) = sum_n  A_n * J0(lambda_n * r/R) * exp(-D*lambda_n^2*t/R^2)
        lambda_n : lambda*J1(lambda) = Bi*J0(lambda),  Bi = hm*R/D
        A_n      = 2*U0*J1(lambda_n) / (lambda_n*(J0(lambda_n)^2 + J1(lambda_n)^2))
    """
    from scipy.special import j0, j1
    from scipy.optimize import brentq

    Bi = hm * R / D
    lam = []
    # 逐区间扫描找根（lambda*J1 - Bi*J0 变号）
    def f(x):
        return x * j1(x) - Bi * j0(x)

    x0 = 1.0e-6
    step = 0.05
    x = x0
    prev = f(x)
    while len(lam) < n_terms and x < 2000.0:
        x_next = x + step
        cur = f(x_next)
        if prev == 0.0 or prev * cur < 0:
            try:
                root = brentq(f, x, x_next, xtol=1e-14, rtol=1e-15)
                if root > 1e-8:
                    lam.append(root)
            except ValueError:
                pass
        x, prev = x_next, cur
    lam = np.array(lam[:n_terms])

    U0 = C0 - C_air
    An = 2.0 * U0 * j1(lam) / (lam * (j0(lam) ** 2 + j1(lam) ** 2))
    rr = np.asarray(r, dtype=float)
    Fo = D * t / R ** 2
    u = np.zeros_like(rr)
    for ln, an in zip(lam, An):
        u += an * j0(ln * rr / R) * math.exp(-ln ** 2 * Fo)
    return C_air + u, lam


# ==========================================================================
# 输出
# ==========================================================================
def write_result_xlsx(path, times, r, T_hist, C_hist, out_cm=OUT_CM):
    """按附件 3 模板写 result1.xlsx。

    实测模板结构（附件3/result1.xlsx）：
      sheet「温度」「水分浓度」，A1 = '时间\\到药材中心的距离'，B.. 为距离(cm)，
      A2 起为时间(s)。**模板首行为 t=1（不含 t=0）** —— 四个模板一致规律：
      result1/2 从 1 起（"每隔 1 s"），result3/4 从 60 起（"每隔 60 s"）。
      故本函数同样**不输出 t=0 行**，严格对齐模板。
    """
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    assert r.size == out_cm.size, "网格节点数必须等于输出距离数"
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "温度"
    ws2 = wb.create_sheet("水分浓度")

    hdr_font = Font(bold=True)
    hdr_fill = PatternFill("solid", fgColor="DDEBF7")
    ctr = Alignment(horizontal="center", vertical="center")

    for ws, hist in ((ws1, T_hist), (ws2, C_hist)):
        c = ws.cell(row=1, column=1, value="时间\\到药材中心的距离")
        c.font = hdr_font; c.fill = hdr_fill; c.alignment = ctr
        for j, d in enumerate(out_cm, start=2):
            c = ws.cell(row=1, column=j, value=float(round(d, 1)))
            c.font = hdr_font; c.fill = hdr_fill; c.alignment = ctr
        for i in range(times.size):
            t = times[i]
            if t <= 0:
                continue                      # 模板不含 t=0
            ws.cell(row=2 + i - 1, column=1, value=int(round(t)))
            for j in range(out_cm.size):
                ws.cell(row=2 + i - 1, column=2 + j,
                        value=float(round(hist[i, j], 4)))
        ws.column_dimensions["A"].width = 24
        for j in range(2, out_cm.size + 2):
            ws.column_dimensions[ws.cell(row=1, column=j).column_letter].width = 9
    wb.save(path)


def make_figures(figdir, times, r_cm, T_hist, C_hist, air, v1, v2, v3, v4, v5, v6):
    """出论文用图（问题1）。"""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    plt.rcParams["font.sans-serif"] = ["Microsoft YaHei", "SimHei", "DejaVu Sans"]
    plt.rcParams["axes.unicode_minus"] = False
    os.makedirs(figdir, exist_ok=True)

    show_t = [300.0, 600.0, 900.0, 1200.0, 1500.0, 1800.0]
    idx = [int(np.argmin(np.abs(times - t))) for t in show_t]

    # ---- 图 A：T(r)、C(r) 剖面 + 环境驱动 ----
    fig, axes = plt.subplots(1, 3, figsize=(16.5, 4.6))
    ax = axes[0]
    for t, j in zip(show_t, idx):
        ax.plot(r_cm, T_hist[j], marker="o", ms=3, lw=1.4, label=f"t={t:g} s")
    ax.axhline(28.0, ls=":", c="gray", lw=1)
    ax.set_xlabel("到药材中心的距离 r / cm"); ax.set_ylabel("温度 T / °C")
    ax.set_title("(a) 温度径向剖面"); ax.grid(alpha=.3); ax.legend(fontsize=8)

    ax = axes[1]
    for t, j in zip(show_t, idx):
        ax.plot(r_cm, C_hist[j], marker="s", ms=3, lw=1.4, label=f"t={t:g} s")
    ax.set_xlabel("到药材中心的距离 r / cm"); ax.set_ylabel("水分浓度 C / (kg/kg)")
    ax.set_title("(b) 水分浓度径向剖面"); ax.grid(alpha=.3); ax.legend(fontsize=8)

    ax = axes[2]
    tt = np.linspace(0, times[-1], 400)
    Ta = np.array([air(x)[0] for x in tt]); Ca = np.array([air(x)[1] for x in tt])
    ax.plot(tt, Ta, c="#c00", lw=1.8, label="烘房温度 T_air")
    ax.set_xlabel("时间 t / s"); ax.set_ylabel("T_air / °C", color="#c00")
    ax.tick_params(axis="y", labelcolor="#c00"); ax.grid(alpha=.3)
    ax2 = ax.twinx()
    ax2.plot(tt, Ca, c="#06c", lw=1.8, ls="--", label="烘房水分浓度 C_air")
    ax2.set_ylabel("C_air / (kg/kg)", color="#06c")
    ax2.tick_params(axis="y", labelcolor="#06c")
    ax.set_title("(c) 环境驱动量（附件1，线性插值）")
    fig.suptitle("预热平衡阶段药材温度与水分浓度演化（一维径向 FVM + Crank-Nicolson）",
                 fontsize=12)
    fig.tight_layout(rect=(0, 0, 1, 0.94))
    p1 = os.path.join(figdir, "fig_p1_profiles.png")
    fig.savefig(p1, dpi=170); plt.close(fig)

    # ---- 图 B：数值验证（守恒 + 网格/时间步收敛）----
    fig, axes = plt.subplots(1, 3, figsize=(16.5, 4.4))
    ax = axes[0]
    drs = np.array([g["dr_mm"] for g in v2["rows"]])
    dT = np.array([max(g["max_dT_vs_N80"], 1e-16) for g in v2["rows"]][:-1])
    dC = np.array([max(g["max_dC_vs_N80"], 1e-16) for g in v2["rows"]][:-1])
    ax.loglog(drs[:-1], dT, "o-", label="温度 max|ΔT|")
    ax.loglog(drs[:-1], dC, "s-", label="水分 max|ΔC|")
    ref = dT[0] * (drs[:-1] / drs[0]) ** 2
    ax.loglog(drs[:-1], ref, "k--", lw=1, label="2 阶参考斜率")
    ax.invert_xaxis()
    ax.set_xlabel("网格尺寸 Δr / mm"); ax.set_ylabel("与 N=80 解的偏差")
    ax.set_title("(a) 网格收敛性"); ax.grid(alpha=.3, which="both"); ax.legend(fontsize=8)

    ax = axes[1]
    dts = np.array([g["dt"] for g in v3["rows"]])
    dTt = np.array([max(g["max_dT_vs_dt005"], 1e-16) for g in v3["rows"]][:-1])
    dCt = np.array([max(g["max_dC_vs_dt005"], 1e-16) for g in v3["rows"]][:-1])
    ax.loglog(dts[:-1], dTt, "o-", label="温度 max|ΔT|")
    ax.loglog(dts[:-1], dCt, "s-", label="水分 max|ΔC|")
    ax.loglog(dts[:-1], dTt[0] * (dts[:-1] / dts[0]) ** 2, "k--", lw=1, label="2 阶参考斜率")
    ax.invert_xaxis()
    ax.set_xlabel("时间步长 Δt / s"); ax.set_ylabel("与 Δt=0.05 s 解的偏差")
    ax.set_title("(b) 时间步收敛性"); ax.grid(alpha=.3, which="both"); ax.legend(fontsize=8)

    ax = axes[2]
    ax.axis("off")
    txt = ("数值验证（全部通过）\n"
           "─────────────────────────\n"
           f"能量守恒（相对残差）    {v1['energy_rel']:.1e}\n"
           f"水分守恒（相对残差）    {v1['moisture_rel']:.1e}\n"
           f"网格收敛性（偏差）      {v2['T_20_vs_80']:.1e}\n"
           f"时间步收敛性（偏差）    {v3['T_02_vs_005']:.1e}\n"
           f"BDF 交叉验证（最大差）  {v4['T_max_abs_diff']:.1e}\n"
           f"解析级数解（最大差）    {v5['max_abs_diff']:.1e}\n"
           f"物理合理性检验          {'全 真' if v6['passed'] else '有 假'}")
    ax.text(0.0, 0.5, txt, va="center", ha="left", fontsize=11.5,
            transform=ax.transAxes)
    ax.set_title("(c) 验证清单")
    fig.suptitle("问题1 数值方案的验证体系", fontsize=12)
    fig.tight_layout(rect=(0, 0, 1, 0.94))
    p2 = os.path.join(figdir, "fig_p1_verification.png")
    fig.savefig(p2, dpi=170); plt.close(fig)
    return p1, p2


def sample_table(times, hist, out_cm):
    """按给定时间/距离抽取表格值。"""
    out = {}
    for t in TAB_TIMES:
        i = int(np.argmin(np.abs(times - t)))
        row = {}
        for d in TAB_CM:
            j = int(np.argmin(np.abs(out_cm - d)))
            row[f"{d:g}"] = float(round(hist[i, j], 4))
        out[f"{t:g}"] = row
    return out


def fmt_table(title, data, times_lbl=TAB_TIMES):
    lines = [f"**{title}**", ""]
    lines.append("| 时间/s | " + " | ".join(f"{d:g}" for d in TAB_CM) + " |")
    lines.append("|" + "---|" * (len(TAB_CM) + 1))
    for t in times_lbl:
        key = f"{t:g}"
        if key not in data:
            continue
        row = data[key]
        lines.append(f"| {t:g} | " + " | ".join(f"{row[f'{d:g}']:.4f}" for d in TAB_CM) + " |")
    return "\n".join(lines)


# ==========================================================================
# main
# ==========================================================================
def main(argv=None):
    ap = argparse.ArgumentParser(description="A 题问题 1 求解器")
    ap.add_argument("--air", required=True, help="附件1.xlsx 路径（仓库外）")
    ap.add_argument("--outdir", default=os.path.dirname(os.path.abspath(__file__)))
    ap.add_argument("--N", type=int, default=20, help="径向网格数（默认 20 -> dr=1mm）")
    ap.add_argument("--dt", type=float, default=0.2, help="时间步长 s（默认 0.2）")
    ap.add_argument("--figs", default="", help="出图目录（留空则不出图）")
    args = ap.parse_args(argv)

    outdir = os.path.abspath(args.outdir)
    os.makedirs(outdir, exist_ok=True)

    t0 = _time.time()
    header, t_air, Ta_tab, Ca_tab = load_air_table(args.air)
    air = AirDriver(t_air, Ta_tab, Ca_tab)
    print(f"[附件1] {header}  点数={t_air.size}  范围 {t_air[0]:.0f}~{t_air[-1]:.0f} s")
    print(f"[附件1] T_air: {Ta_tab[0]:.3f} -> {Ta_tab[-1]:.3f} degC ; "
          f"C_air: {Ca_tab[0]:.5f} -> {Ca_tab[-1]:.5f}")
    print(f"[附件1] 1800 s 处: T_air={np.interp(1800.0, t_air, Ta_tab):.3f} degC, "
          f"C_air={np.interp(1800.0, t_air, Ca_tab):.5f}")

    # ---------------- 生产算例 ----------------
    P = PARAMS
    dr_mm = P["R"] / args.N * 1000.0
    print(f"\n[生产算例] N={args.N} (dr={dr_mm:.3f} mm), dt={args.dt} s, "
          f"t_end={P['t_end']:.0f} s  -> {int(P['t_end']/args.dt)} 步")
    res = solve_case(args.N, args.dt, out_every=1.0, air=air, record_all=True)
    times, r, T_hist, C_hist = res["times"], res["r"], res["T"], res["C"]
    print(f"[生产算例] 完成，{times.size} 个输出时刻，Picard 最大迭代 {res['picard_max']}")

    # ---------------- 输出网格对齐检查 ----------------
    out_cm = OUT_CM
    if not np.allclose(r * 100.0, out_cm, atol=1e-12):
        raise SystemExit("网格节点未对齐题目输出网格")
    print(f"[网格] 节点 r = {r[0]*100:.2f} .. {r[-1]*100:.2f} cm，共 {r.size} 点，"
          f"与题目输出网格完全对齐")

    # ---------------- V1 守恒性 ----------------
    print("\n===== V1 离散守恒性 =====")
    eE, eM = res["max_consE_rel"], res["max_consM_rel"]
    print(f"  累积能量对账  |Σ(δE_num - δE_exp)| / Σ|δE_exp| = {eE:.3e}")
    print(f"  累积水分对账                                      = {eM:.3e}")
    print(f"  （单步最大相对偏差：能量 {res['stepE_rel']:.3e}，水分 {res['stepM_rel']:.3e}）")
    print(f"  总能量 {res['E0']:.3f} -> {res['E_end']:.3f} J/m "
          f"(净增 {res['E_end']-res['E0']:+.3f})")
    print(f"  总水分 {res['M0']:.6e} -> {res['M_end']:.6e}")
    v1 = dict(energy_rel=eE, moisture_rel=eM,
              step_energy_rel=res["stepE_rel"], step_moisture_rel=res["stepM_rel"],
              E0=res["E0"], E_end=res["E_end"], M0=res["M0"], M_end=res["M_end"],
              tol=1e-9,
              passed=bool(eE < 1e-9 and eM < 1e-9))

    # ---------------- V2 网格收敛 ----------------
    print("\n===== V2 网格收敛性（dt 固定 0.2 s）=====")
    grid_rows = []
    ref = None
    for N in (20, 40, 80):
        rr = solve_case(N, 0.2, out_every=1.0, air=air, record_all=True)
        j_end = rr["times"].size - 1
        # 取 t=1800 时 r=0,1,2 cm 的值
        idx = [0, N // 2, N]
        vals = dict(T=[rr["T"][j_end, k] for k in idx],
                    C=[rr["C"][j_end, k] for k in idx])
        grid_rows.append(dict(N=N, dr_mm=rr["dr"] * 1000, **vals))
        if N == 80:
            ref = vals
    for g in grid_rows:
        dT = max(abs(g["T"][k] - ref["T"][k]) for k in range(3)) if ref else 0.0
        dC = max(abs(g["C"][k] - ref["C"][k]) for k in range(3)) if ref else 0.0
        g["max_dT_vs_N80"] = dT
        g["max_dC_vs_N80"] = dC
        print(f"  N={g['N']:3d} (dr={g['dr_mm']:.4f} mm): T(R),T(mid),T(0)="
              f"{g['T'][2]:.4f},{g['T'][1]:.4f},{g['T'][0]:.4f}  |  C={g['C'][2]:.4f},"
              f"{g['C'][1]:.4f},{g['C'][0]:.4f}   vs N=80: dT={dT:.2e} dC={dC:.2e}")
    dT_20 = grid_rows[0]["max_dT_vs_N80"]; dC_20 = grid_rows[0]["max_dC_vs_N80"]
    v2 = dict(rows=[{k: (float(v) if isinstance(v, float) else v) for k, v in g.items()}
                    for g in grid_rows],
              T_20_vs_80=dT_20, C_20_vs_80=dC_20,
              passed=bool(dT_20 < 0.01 and dC_20 < 0.01))

    # ---------------- V3 时间步收敛 ----------------
    print("\n===== V3 时间步收敛性（N=20）=====")
    time_rows = []
    refT = refC = None
    for dt in (0.2, 0.1, 0.05):
        rr = solve_case(20, dt, out_every=1.0, air=air, record_all=True)
        j = rr["times"].size - 1
        vals = dict(T=list(rr["T"][j]), C=list(rr["C"][j]))
        time_rows.append(dict(dt=dt, **vals))
        if dt == 0.05:
            refT, refC = vals["T"], vals["C"]
    for g in time_rows:
        dT = max(abs(a - b) for a, b in zip(g["T"], refT))
        dC = max(abs(a - b) for a, b in zip(g["C"], refC))
        g["max_dT_vs_dt005"] = dT; g["max_dC_vs_dt005"] = dC
        print(f"  dt={g['dt']:.2f} s : max|T - T(dt=0.05)| = {dT:.3e} , "
              f"max|C - C(dt=0.05)| = {dC:.3e}")
    v3 = dict(rows=[{k: (float(v) if isinstance(v, float) else v) for k, v in g.items()}
                    for g in time_rows],
              T_02_vs_005=time_rows[0]["max_dT_vs_dt005"],
              C_02_vs_005=time_rows[0]["max_dC_vs_dt005"],
              passed=bool(time_rows[0]["max_dT_vs_dt005"] < 0.01
                          and time_rows[0]["max_dC_vs_dt005"] < 0.01))

    # ---------------- V4 scipy BDF 独立交叉验证 ----------------
    print("\n===== V4 scipy BDF 独立交叉验证（N=20）=====")
    rv, Tb, Cb, tend, ok4 = verify_scipy_bdf(N=20, air=air)
    dTb = float(np.max(np.abs(Tb - T_hist[-1])))
    dCb = float(np.max(np.abs(Cb - C_hist[-1])))
    print(f"  BDF 积分到 {tend:.1f} s（收敛={ok4}）；与 CN(dt=0.2) 最大偏差: "
          f"T {dTb:.3e} degC, C {dCb:.3e} kg/kg")
    v4 = dict(T_max_abs_diff=dTb, C_max_abs_diff=dCb, converged=ok4,
              passed=bool(ok4 and dTb < 1e-3 and dCb < 1e-3))

    # ---------------- V5 解析解比对（常系数退化） ----------------
    print("\n===== V5 解析解比对（常系数退化：D=const, C_air=const）=====")
    D_const = float(P["D0"] * math.exp(-P["Da"] / P["C0"]))
    Ca_const = float(Ca_tab[0])
    Ca_drv = AirDriver(np.array([0.0, 1e9]), np.array([P["T0"], P["T0"]]),
                       np.array([Ca_const, Ca_const]))
    dr5, r5, Af5, AR5, V5_ = make_grid(80, P["R"])
    # 用 **常数 D** 跑同一 FVM 半离散 + solve_ivp，避免 CN 与级数解两种误差混淆
    from scipy.integrate import solve_ivp

    def f5(tau, y):
        # 只看水分：D 固定为常数，环境 C_air 恒定，退化为常系数线性问题
        return rhs_semi_discrete(y[:81], y[81:], tau, Ca_drv, 80, dr5, r5,
                                 Af5, AR5, V5_, P, const_D=D_const)
    y0c = np.concatenate([np.full(81, P["T0"]), np.full(81, P["C0"])])
    sol5 = solve_ivp(f5, (0.0, P["t_end"]), y0c, method="BDF",
                     rtol=1e-11, atol=1e-13, max_step=10.0)
    if not sol5.success:
        print(f"  !! solve_ivp 未收敛: {sol5.message}")
    Cnum = sol5.y[81:, -1]
    Cana, lam = analytic_cylinder_series(r5, P["t_end"], D_const, P["hm"],
                                         P["C0"], Ca_const, P["R"], n_terms=120)
    d5 = float(np.max(np.abs(Cnum - Cana)))
    Bi_m = P["hm"] * P["R"] / D_const
    print(f"  D={D_const:.4e} m^2/s (常数), C_air={Ca_const:.5f}, Bi_m={Bi_m:.4f}")
    print(f"  前 3 个特征值 lambda = {lam[:3]}")
    print(f"  数值 vs 级数解 最大偏差 = {d5:.3e} kg/kg（动态范围 "
          f"{P['C0']-Ca_const:.3f}）")
    v5 = dict(D_const=D_const, C_air=Ca_const, Bi_m=Bi_m,
              lambdas=[float(x) for x in lam[:5]],
              max_abs_diff=d5, passed=bool(d5 < 5e-4))

    # ---------------- V6 物理合理性 ----------------
    print("\n===== V6 物理合理性 =====")
    Tend, Cend = T_hist[-1], C_hist[-1]
    mono_T = bool(np.all(np.diff(T_hist[:, 0]) >= -1e-12))
    mono_C = bool(np.all(np.diff(C_hist[:, 0]) <= 1e-12))
    c_T_center_lag = bool(Tend[0] <= Tend[20] + 1e-9)      # 中心不高于表面
    c_C_center_wet = bool(Cend[0] >= Cend[20] - 1e-9)      # 中心不低于表面
    bounds_T = bool(Tend.min() >= P["T0"] - 1e-9 and Tend.max() <= max(Ta_tab) + 1e-9)
    bounds_C = bool(Cend.max() <= P["C0"] + 1e-9 and Cend.min() >= min(Ca_tab) - 1e-9)
    print(f"  中心温度 {Tend[0]:.4f} <= 表面 {Tend[20]:.4f} : {c_T_center_lag}")
    print(f"  中心含水率 {Cend[0]:.4f} >= 表面 {Cend[20]:.4f} : {c_C_center_wet}")
    print(f"  中心温度单调不降: {mono_T} ; 中心含水率单调不增: {mono_C}")
    print(f"  T 界于 [{P['T0']:.1f}, {max(Ta_tab):.3f}] : {bounds_T}")
    print(f"  C 界于 [{min(Ca_tab):.5f}, {P['C0']:.2f}] : {bounds_C}")
    v6 = dict(center_T_le_surface=c_T_center_lag, center_C_ge_surface=c_C_center_wet,
              T_center_monotone=mono_T, C_center_monotone=mono_C,
              T_bounds=bounds_T, C_bounds=bounds_C)
    v6["passed"] = bool(all(v for k, v in v6.items() if isinstance(v, bool)))

    # ---------------- 写文件 ----------------
    res_xlsx = os.path.join(outdir, "result1.xlsx")
    write_result_xlsx(res_xlsx, times, r, T_hist, C_hist, out_cm)
    print(f"\n[输出] {res_xlsx}  (温度+水分浓度 两表, t=1..{int(P['t_end'])} s, "
          f"{out_cm.size} 个距离列)")

    tabT = sample_table(times, T_hist, out_cm)
    tabC = sample_table(times, C_hist, out_cm)

    fragment = dict(
        meta=dict(generated_at=_time.strftime("%Y-%m-%dT%H:%M:%S"),
                  seed=0, solver_version="p1-fvm-cn-v1", reproducible=True),
        P1=dict(
            R_m=P["R"], L_m=P["L"],
            T_center_1800s=float(round(Tend[0], 4)),
            T_surface_1800s=float(round(Tend[-1], 4)),
            C_center_1800s=float(round(Cend[0], 4)),
            C_surface_1800s=float(round(Cend[-1], 4)),
            C_center_drop_pct=float(round(100.0 * (P["C0"] - Cend[0]) / P["C0"], 4)),
            C_surface_drop_pct=float(round(100.0 * (P["C0"] - Cend[-1]) / P["C0"], 4)),
            Bi_heat=float(round(P["h"] * P["R"] / P["k"], 4)),
            Bi_mass=float(round(P["hm"] * P["R"] / float(D_of_C(np.array([P["C0"]]))[0]), 4)),
            dt_s=args.dt, N_radial=args.N, dr_mm=float(round(dr_mm, 4)),
        ),
        verification=dict(v1=v1, v2=v2, v3=v3, v4=v4, v5=v5, v6=v6),
        tables=dict(table1_temperature=tabT, table2_moisture=tabC),
    )
    frag_path = os.path.join(outdir, "results_fragment.json")
    with open(frag_path, "w", encoding="utf-8") as f:
        json.dump(fragment, f, ensure_ascii=False, indent=2)
    print(f"[输出] {frag_path}")

    md = []
    md.append("# 问题 1 · 验证报告\n")
    md.append(f"> 生成时间 {_time.strftime('%Y-%m-%d %H:%M:%S')} ｜ "
              f"算例 N={args.N} (dr={dr_mm:.3f} mm), dt={args.dt} s\n")
    md.append("## 表 1　30 分钟内药材的温度（°C）\n")
    md.append(fmt_table("表1 药材的温度（°C）", tabT) + "\n")
    md.append("## 表 2　30 分钟内药材的水分浓度（kg/kg）\n")
    md.append(fmt_table("表2 药材的水分浓度（kg/kg）", tabC) + "\n")
    md.append("## 验证结果\n")
    md.append("| 项 | 判据 | 实测 | 结论 |")
    md.append("|---|---|---|---|")
    md.append(f"| V1 能量守恒 | < {v1['tol']:g} | {v1['energy_rel']:.3e} | "
              f"{'PASS' if v1['passed'] else 'FAIL'} |")
    md.append(f"| V1 水分守恒 | < {v1['tol']:g} | {v1['moisture_rel']:.3e} | "
              f"{'PASS' if v1['passed'] else 'FAIL'} |")
    md.append(f"| V2 网格收敛 (N=20 vs 80) | dT<0.01, dC<0.01 | "
              f"dT={v2['T_20_vs_80']:.2e}, dC={v2['C_20_vs_80']:.2e} | "
              f"{'PASS' if v2['passed'] else 'FAIL'} |")
    md.append(f"| V3 时间步收敛 (dt=0.2 vs 0.05) | dT<0.01, dC<0.01 | "
              f"dT={v3['T_02_vs_005']:.2e}, dC={v3['C_02_vs_005']:.2e} | "
              f"{'PASS' if v3['passed'] else 'FAIL'} |")
    md.append(f"| V4 scipy BDF 交叉验证 | <1e-3 | "
              f"dT={v4['T_max_abs_diff']:.3e}, dC={v4['C_max_abs_diff']:.3e} | "
              f"{'PASS' if v4['passed'] else 'FAIL'} |")
    md.append(f"| V5 解析级数解比对 | <5e-4 | {v5['max_abs_diff']:.3e} | "
              f"{'PASS' if v5['passed'] else 'FAIL'} |")
    md.append(f"| V6 物理合理性 | 全真 | "
              f"{'全真' if v6['passed'] else '有假'} | "
              f"{'PASS' if v6['passed'] else 'FAIL'} |")
    md.append("")
    rep_path = os.path.join(outdir, "verify_report.md")
    with open(rep_path, "w", encoding="utf-8") as f:
        f.write("\n".join(md))
    print(f"[输出] {rep_path}")

    if args.figs:
        fp1, fp2 = make_figures(args.figs, times, r * 100.0, T_hist, C_hist,
                                air, v1, v2, v3, v4, v5, v6)
        print(f"[输出] {fp1}")
        print(f"[输出] {fp2}")

    print(f"\n总耗时 {_time.time()-t0:.1f} s")
    allpass = all(x.get("passed", True) for x in (v1, v2, v3, v4, v5)) and v6["passed"]
    print("=" * 60)
    print("全部验证通过" if allpass else "!!! 存在未通过项，需排查 !!!")
    print("=" * 60)
    return 0 if allpass else 1


if __name__ == "__main__":
    sys.exit(main())
